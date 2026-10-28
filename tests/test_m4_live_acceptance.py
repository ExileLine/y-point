# -*- coding: utf-8 -*-
# @Time    : 2026-10-28 14:45:00
# @Author  : yangyuexiong
# @File    : test_m4_live_acceptance.py

"""通过 Ollama 真实模型执行可选的 M4 端到端验收。"""

import asyncio
from io import BytesIO
import os
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.config import get_config
from app.core.password import hash_password
from app.db.session import get_db_session
from app.main import app
from app.models.admin import Admin
from app.models.base import Base
from app.models.model_catalog import ModelCatalog
from app.models.provider_connection import ProviderConnection
from app.services.run_dispatcher import dispatch_next_run
from app.services.worker_tools import make_worker_tool_resolver


@pytest.mark.skipif(os.environ.get("YPOINT_TEST_OLLAMA_M4") != "1", reason="Requires local Ollama")
def test_live_m4_document_generation(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """真实模型生成的文件随助手消息返回并可授权下载。"""

    async def exercise() -> None:
        """在隔离库中验证上传、多格式生成、多附件和旧产物续编辑。"""
        engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'm4-live.db'}")
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        monkeypatch.setattr(get_config(), "ARTIFACT_STORAGE_DIR", str(tmp_path / "artifacts"))

        async def test_session():
            """将 HTTP 请求绑定到隔离业务库。"""
            async with sessions() as db:
                yield db

        app.dependency_overrides[get_db_session] = test_session
        try:
            async with engine.begin() as connection:
                await connection.run_sync(Base.metadata.create_all)
            async with sessions.begin() as db:
                db.add_all(
                    [
                        Admin(id=1, username="m4-live", password=hash_password("StrongPassword123"), role="user", status=1, is_deleted=0),
                        ProviderConnection(
                            id=1,
                            name="m4-ollama",
                            name_key="m4-ollama",
                            kind="ollama",
                            base_url="http://127.0.0.1:11434/v1",
                            timeout_seconds=120,
                            creator_id=1,
                            status=1,
                            is_deleted=0,
                        ),
                        ModelCatalog(
                            id=1,
                            provider_id=1,
                            external_id="gemma4:31b-cloud",
                            display_name="Gemma4",
                            capabilities=["text", "tools"],
                            verified_capabilities=["text", "tools"],
                            capabilities_verified=True,
                            context_window=32768,
                            creator_id=1,
                            status=1,
                            is_deleted=0,
                        ),
                    ]
                )

            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                login = await client.post(
                    "/api/v1/auth/login",
                    json={
                        "identifier": "m4-live",
                        "password": "StrongPassword123",
                    },
                )
                assert login.status_code == 200, login.text
                headers = {
                    "Authorization": f"Bearer {login.json()['data']['access_token']}",
                }
                created = await client.post(
                    "/api/v1/agents",
                    headers=headers,
                    json={
                        "name": "M4 real acceptance",
                        "default_model_id": 1,
                        "system_prompt": "When asked to create a file, call the matching file creation tool. "
                        "Use the supplied reference facts. Keep the file concise.",
                        "max_output_tokens": 1024,
                        "max_model_calls": 6,
                        "max_tool_calls": 4,
                    },
                )
                assert created.status_code == 201, created.text
                base = f"/api/v1/agents/{created.json()['data']['id']}/sessions"
                session = await client.post(
                    base,
                    headers=headers,
                    json={
                        "title": "M4 real files",
                    },
                )
                assert session.status_code == 201, session.text
                session_base = f"{base}/{session.json()['data']['id']}"
                upload = await client.post(
                    "/api/v1/artifacts/uploads",
                    headers=headers,
                    files={
                        "files": ("facts.txt", "Project ALPHA-42. Budget 17. Owner Alice.".encode(), "text/plain"),
                    },
                )
                assert upload.status_code == 201, upload.text
                source_id = upload.json()["data"][0]["id"]

                first_document_id = None
                first_pdf_id = None
                for suffix, instruction in (
                    ("docx", "Create a DOCX document summarizing the attached facts."),
                    ("pdf", "Create a PDF document summarizing the attached facts."),
                    ("xlsx", "Create an XLSX spreadsheet with the attached facts as rows."),
                ):
                    submitted = await client.post(
                        f"{session_base}/runs",
                        headers=headers,
                        json={
                            "request_key": f"m4-live-{suffix}",
                            "input_text": instruction,
                            "artifact_ids": [source_id],
                            "context_artifact_ids": [source_id],
                        },
                    )
                    assert submitted.status_code == 201, submitted.text
                    run_id = submitted.json()["data"]["id"]
                    result = await dispatch_next_run(
                        sessions,
                        tmp_path / "checkpoints.db",
                        timeout_seconds=120,
                        lease_seconds=30,
                        poll_seconds=2,
                        resolve_tools=make_worker_tool_resolver(sessions),
                    )
                    assert result is not None and result.id == run_id
                    assert result.state == "completed", (suffix, result.state, result.failure_code)
                    messages = await client.get(f"{session_base}/messages", headers=headers)
                    assert messages.status_code == 200, messages.text
                    assistant = next((item for item in messages.json()["data"] if item["run_id"] == run_id and item["role"] == "assistant"))
                    matches = [item for item in assistant["artifacts"] if item["filename"].endswith(f".{suffix}")]
                    assert len(matches) == 1 and matches[0]["available"], (suffix, assistant)
                    artifact = matches[0]
                    assert artifact["context_artifact_ids"] == [source_id]
                    if suffix == "docx":
                        first_document_id = artifact["artifact_id"]
                    if suffix == "pdf":
                        first_pdf_id = artifact["artifact_id"]
                    download = await client.get(f"/api/v1/artifacts/{artifact['artifact_id']}/download", headers=headers)
                    assert download.status_code == 200
                    if suffix == "docx":
                        content = "\n".join(paragraph.text for paragraph in Document(BytesIO(download.content)).paragraphs)
                        assert "ALPHA-42" in content
                    elif suffix == "pdf":
                        pages = PdfReader(BytesIO(download.content)).pages
                        assert pages and "ALPHA-42" in "\n".join(page.extract_text() for page in pages)
                    else:
                        workbook = load_workbook(BytesIO(download.content), read_only=True)
                        try:
                            assert workbook.sheetnames
                            values = [str(value) for sheet in workbook for row in sheet.values for value in row if value is not None]
                            assert "ALPHA-42" in " ".join(values)
                        finally:
                            workbook.close()

                assert first_document_id is not None
                revised = await client.post(
                    f"{session_base}/runs",
                    headers=headers,
                    json={
                        "request_key": "m4-live-revision",
                        "input_text": f"Read the existing DOCX artifact {first_document_id}, then create "
                        "a revised DOCX including the original facts and 'Approved on Tuesday'.",
                    },
                )
                assert revised.status_code == 201, revised.text
                revision_id = revised.json()["data"]["id"]
                result = await dispatch_next_run(
                    sessions,
                    tmp_path / "checkpoints.db",
                    timeout_seconds=120,
                    lease_seconds=30,
                    poll_seconds=2,
                    resolve_tools=make_worker_tool_resolver(sessions),
                )
                assert result is not None and result.id == revision_id
                assert result.state == "completed", (result.state, result.failure_code)
                messages = (await client.get(f"{session_base}/messages", headers=headers)).json()["data"]
                assistant = next(item for item in messages if item["run_id"] == revision_id and item["role"] == "assistant")
                revised_files = [item for item in assistant["artifacts"] if item["filename"].endswith(".docx")]
                assert len(revised_files) == 1
                assert revised_files[0]["read_source_artifact_ids"] == [first_document_id]
                download = await client.get(f"/api/v1/artifacts/{revised_files[0]['artifact_id']}/download", headers=headers)
                assert download.status_code == 200
                document = Document(BytesIO(download.content))
                content = "\n".join(paragraph.text for paragraph in document.paragraphs)
                assert "ALPHA-42" in content and "Approved on Tuesday" in content

                assert first_pdf_id is not None
                revised_pdf = await client.post(
                    f"{session_base}/runs",
                    headers=headers,
                    json={
                        "request_key": "m4-live-pdf-revision",
                        "input_text": "Read the PDF you created earlier, then create a revised PDF that includes "
                        "the original project code and 'Reviewed on Friday'.",
                    },
                )
                assert revised_pdf.status_code == 201, revised_pdf.text
                pdf_run_id = revised_pdf.json()["data"]["id"]
                result = await dispatch_next_run(
                    sessions,
                    tmp_path / "checkpoints.db",
                    timeout_seconds=120,
                    lease_seconds=30,
                    poll_seconds=2,
                    resolve_tools=make_worker_tool_resolver(sessions),
                )
                assert result is not None and result.id == pdf_run_id
                assert result.state == "completed", (result.state, result.failure_code)
                messages = (await client.get(f"{session_base}/messages", headers=headers)).json()["data"]
                assistant = next(item for item in messages if item["run_id"] == pdf_run_id and item["role"] == "assistant")
                pdf_files = [item for item in assistant["artifacts"] if item["filename"].endswith(".pdf")]
                assert len(pdf_files) == 1
                assert pdf_files[0]["read_source_artifact_ids"] == [first_pdf_id]
                download = await client.get(f"/api/v1/artifacts/{pdf_files[0]['artifact_id']}/download", headers=headers)
                assert download.status_code == 200
                pdf_text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(download.content)).pages)
                assert "ALPHA-42" in pdf_text and "Reviewed on Friday" in pdf_text

                second_upload = await client.post(
                    "/api/v1/artifacts/uploads",
                    headers=headers,
                    files={
                        "files": ("second.txt", b"Project BETA-73. Budget 29. Owner Bob.", "text/plain"),
                    },
                )
                assert second_upload.status_code == 201, second_upload.text
                second_id = second_upload.json()["data"][0]["id"]
                combined = await client.post(
                    f"{session_base}/runs",
                    headers=headers,
                    json={
                        "request_key": "m4-live-two-attachments",
                        "input_text": "Create a TXT file listing the project codes from both attached files.",
                        "artifact_ids": [source_id, second_id],
                        "context_artifact_ids": [source_id, second_id],
                    },
                )
                assert combined.status_code == 201, combined.text
                combined_id = combined.json()["data"]["id"]
                result = await dispatch_next_run(
                    sessions,
                    tmp_path / "checkpoints.db",
                    timeout_seconds=120,
                    lease_seconds=30,
                    poll_seconds=2,
                    resolve_tools=make_worker_tool_resolver(sessions),
                )
                assert result is not None and result.id == combined_id
                assert result.state == "completed", (result.state, result.failure_code)
                messages = (await client.get(f"{session_base}/messages", headers=headers)).json()["data"]
                assistant = next(item for item in messages if item["run_id"] == combined_id and item["role"] == "assistant")
                text_files = [item for item in assistant["artifacts"] if item["filename"].endswith(".txt")]
                assert len(text_files) == 1
                assert text_files[0]["context_artifact_ids"] == [source_id, second_id]
                download = await client.get(f"/api/v1/artifacts/{text_files[0]['artifact_id']}/download", headers=headers)
                assert download.status_code == 200
                assert "ALPHA-42" in download.text and "BETA-73" in download.text
        finally:
            app.dependency_overrides.pop(get_db_session, None)
            await engine.dispose()

    asyncio.run(exercise())
