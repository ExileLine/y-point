"""验证生成文件是真实格式，且 HTML 和表格内容不会变成可执行代码。"""

from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.artifact_renderers import (
    ArtifactSheet,
    render_document_artifact,
    render_spreadsheet_artifact,
    render_text_artifact,
)


def test_document_formats_can_be_reopened() -> None:
    content = "# 夷陵之战\n\n蜀汉与东吴交战。\n\n- 时间：222年"
    document = Document(BytesIO(render_document_artifact("summary.docx", content)))
    assert any("夷陵之战" in paragraph.text for paragraph in document.paragraphs)
    assert any("蜀汉与东吴" in paragraph.text for paragraph in document.paragraphs)
    pdf = PdfReader(BytesIO(render_document_artifact("summary.pdf", content)))
    assert len(pdf.pages) >= 1
    assert pdf.pages[0].mediabox.width > 0


def test_spreadsheet_preserves_text_and_multiple_sheets() -> None:
    sheets = [
        ArtifactSheet(name="概览", rows=[["项目", "数值"], ["结果", 42], ["用户输入", "=HYPERLINK(\"bad\")"]]),
        ArtifactSheet(name="明细", rows=[["项目", "状态"], ["A", True]]),
    ]
    workbook = load_workbook(BytesIO(render_spreadsheet_artifact(sheets)), read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["概览", "明细"]
        assert workbook["概览"]["B2"].value == 42
        assert workbook["概览"]["B3"].value == '=HYPERLINK("bad")'
        assert workbook["概览"]["B3"].data_type == "s"
    finally:
        workbook.close()
    with pytest.raises(ValueError, match="重复"):
        render_spreadsheet_artifact([
            ArtifactSheet(name="Same", rows=[["a"]]), ArtifactSheet(name="same", rows=[["b"]]),
        ])


def test_markdown_and_html_are_utf8_and_html_is_sanitized() -> None:
    assert render_text_artifact("summary.md", "# 标题").decode() == "# 标题"
    content = render_text_artifact("summary.html", "# 标题\n<script>alert(1)</script>\n[link](javascript:alert(1))")
    assert "<h1>标题</h1>" in content.decode()
    assert "<script" not in content.decode()
    assert "javascript:" not in content.decode()
    assert render_text_artifact("data.json", '{"name":"夷陵"}').decode().startswith('{\n  "name": "夷陵"')
    with pytest.raises(ValueError, match="JSON"):
        render_text_artifact("data.json", "{invalid}")
    csv_text = render_text_artifact("data.csv", 'name,value\nrow,"=1+1"').decode()
    assert "'=1+1" in csv_text
