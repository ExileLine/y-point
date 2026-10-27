"""从已授权的上传文件中提取有界的模型参考文本。"""

import io
from html.parser import HTMLParser

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_CONTEXT_ATTACHMENT_BYTES = 16 * 1024
CONTEXT_MEDIA_TYPES = frozenset({
    "text/plain", "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
})
GENERATED_CONTEXT_MEDIA_TYPES = CONTEXT_MEDIA_TYPES | frozenset({
    "text/markdown", "text/html", "text/csv", "application/json",
})


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"script", "style"}:
            self.skip += 1
        elif tag in {"p", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.lines.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self.skip:
            self.skip -= 1
        elif tag in {"p", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.lines.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skip:
            self.lines.append(data)


def _bounded_text(lines) -> str:
    collected = []
    size = 0
    for index, line in enumerate(lines):
        if index >= 5000:
            raise ValueError("文档内容超过行数上限")
        line = line.strip()
        if not line:
            continue
        size += len(line.encode("utf-8")) + (1 if collected else 0)
        if size > MAX_CONTEXT_ATTACHMENT_BYTES:
            raise ValueError("文档提取文本超过 16 KiB 上限")
        collected.append(line)
    if not collected:
        raise ValueError("文档没有可提取的文本")
    return "\n".join(collected)


def extract_context_text(content: bytes, media_type: str) -> str:
    """只提取小型文本上下文；不默默截断未读完的文档。"""
    if media_type in {"text/plain", "text/markdown", "text/csv", "application/json"}:
        return _bounded_text(content.decode("utf-8").splitlines())
    if media_type == "text/html":
        parser = _HTMLText()
        parser.feed(content.decode("utf-8"))
        return _bounded_text("".join(parser.lines).splitlines())
    if media_type == "application/pdf":
        try:
            reader = PdfReader(io.BytesIO(content), strict=True)
        except PdfReadError as exc:
            raise ValueError("PDF 文件无法解析") from exc
        if reader.is_encrypted or not 1 <= len(reader.pages) <= 50:
            raise ValueError("PDF 已加密、为空或超过 50 页")
        try:
            return _bounded_text(line for page in reader.pages for line in (page.extract_text() or "").splitlines())
        except PdfReadError as exc:
            raise ValueError("PDF 文件无法提取文本") from exc
    if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        document = Document(io.BytesIO(content))

        def document_lines():
            for block in document.iter_inner_content():
                if isinstance(block, Paragraph):
                    yield block.text
                elif isinstance(block, Table):
                    for row in block.rows:
                        yield "\t".join(cell.text for cell in row.cells)

        return _bounded_text(document_lines())
    if media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            if not 1 <= len(workbook.worksheets) <= 5:
                raise ValueError("工作表数量超过限制")

            def workbook_lines():
                for sheet in workbook.worksheets:
                    if sheet.max_row > 1000 or sheet.max_column > 100:
                        raise ValueError("工作表行列数超过限制")
                    yield f"[工作表: {sheet.title}]"
                    for row in sheet.iter_rows(values_only=True):
                        yield "\t".join("" if value is None else str(value) for value in row)

            return _bounded_text(workbook_lines())
        finally:
            workbook.close()
    raise ValueError("附件格式不能作为模型文本上下文")
