"""将模型提供的结构化内容转换为可校验的办公文件。"""

import csv
import html
import io
import json
import math
import re
from pathlib import Path

import bleach
import markdown
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import IllegalCharacterError
from pydantic import BaseModel, Field
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

MAX_SOURCE_CHARS = 100_000
MAX_SHEETS = 5
MAX_ROWS_PER_SHEET = 1_000
MAX_COLUMNS = 100
MAX_CELLS = 10_000


class ArtifactSheet(BaseModel):
    """一个工作表的名称和按行排列的单元格。"""

    name: str = Field(min_length=1, max_length=31)
    rows: list[list[str | int | float | bool | None]]


def _lines(content: str):
    for line in content.splitlines():
        if heading := re.match(r"^(#{1,6})\s+(.+)$", line):
            yield "heading", len(heading.group(1)), heading.group(2)
        elif item := re.match(r"^\s*[-*+]\s+(.+)$", line):
            yield "bullet", 0, item.group(1)
        elif item := re.match(r"^\s*\d+[.)]\s+(.+)$", line):
            yield "number", 0, item.group(1)
        elif line.strip():
            yield "paragraph", 0, line.strip()
        else:
            yield "space", 0, ""


def render_text_artifact(filename: str, content: str) -> bytes:
    """生成 UTF-8 文本、结构化数据或经过白名单清洗的 HTML。"""
    if not content.strip() or len(content) > MAX_SOURCE_CHARS:
        raise ValueError("生成内容为空或超过长度限制")
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return content.encode("utf-8")
    if suffix == ".json":
        try:
            parsed = json.loads(content)
        except json.JSONDecodeError as exc:
            raise ValueError("生成内容不是有效 JSON") from exc
        return json.dumps(parsed, ensure_ascii=False, indent=2, allow_nan=False).encode("utf-8")
    if suffix == ".csv":
        try:
            rows = list(csv.reader(io.StringIO(content, newline=""), strict=True))
        except csv.Error as exc:
            raise ValueError("生成内容不是有效 CSV") from exc
        if not rows or len(rows) > MAX_ROWS_PER_SHEET or sum(len(row) for row in rows) > MAX_CELLS:
            raise ValueError("CSV 行数或单元格数量超过限制")
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        for row in rows:
            if len(row) > MAX_COLUMNS:
                raise ValueError("CSV 列数超过限制")
            writer.writerow(["'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value for value in row])
        return output.getvalue().encode("utf-8")
    if suffix != ".html":
        raise ValueError("文本工具仅支持 TXT、Markdown、HTML、JSON 和 CSV")
    body = markdown.markdown(content, extensions=["tables", "fenced_code"])
    body = bleach.clean(
        body,
        tags={"p", "br", "h1", "h2", "h3", "h4", "h5", "h6", "strong", "em", "blockquote",
              "ul", "ol", "li", "pre", "code", "table", "thead", "tbody", "tr", "th", "td", "a", "hr"},
        attributes={"a": ["href", "title"]}, protocols={"http", "https", "mailto"}, strip=True,
    )
    title = html.escape(Path(filename).stem, quote=True)
    return (
        "<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
        f"<title>{title}</title></head><body>{body}</body></html>"
    ).encode()


def render_document_artifact(filename: str, content: str) -> bytes:
    """将 Markdown 风格的标题、段落和列表生成 DOCX 或 PDF。"""
    if not content.strip() or len(content) > MAX_SOURCE_CHARS:
        raise ValueError("生成内容为空或超过长度限制")
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        document = Document()
        for kind, level, value in _lines(content):
            if kind == "heading":
                document.add_heading(value, level=min(level, 4))
            elif kind in {"bullet", "number"}:
                document.add_paragraph(value, style="List Bullet" if kind == "bullet" else "List Number")
            elif kind == "paragraph":
                document.add_paragraph(value)
        output = io.BytesIO()
        document.save(output)
        data = output.getvalue()
        Document(io.BytesIO(data))
        return data
    if suffix != ".pdf":
        raise ValueError("文档工具仅支持 DOCX 和 PDF")
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = {
        "paragraph": ParagraphStyle("body", fontName="STSong-Light", fontSize=11, leading=17, spaceAfter=8),
        "heading": ParagraphStyle("heading", fontName="STSong-Light", fontSize=16, leading=22, spaceBefore=12, spaceAfter=8),
        "list": ParagraphStyle("list", fontName="STSong-Light", fontSize=11, leading=17, leftIndent=18, spaceAfter=4),
    }
    story = []
    for kind, _, value in _lines(content):
        if kind == "space":
            story.append(Spacer(1, 6))
        else:
            prefix = "&#8226; " if kind == "bullet" else ""
            style = styles["heading" if kind == "heading" else "list" if kind in {"bullet", "number"} else "paragraph"]
            story.append(Paragraph(prefix + html.escape(value), style))
    output = io.BytesIO()
    SimpleDocTemplate(output, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48,
                      title=Path(filename).stem).build(story)
    data = output.getvalue()
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted or not reader.pages:
        raise ValueError("PDF 生成后校验失败")
    return data


def render_spreadsheet_artifact(sheets: list[ArtifactSheet]) -> bytes:
    """生成无公式的 XLSX，多表按输入顺序保存。"""
    if not 1 <= len(sheets) <= MAX_SHEETS:
        raise ValueError("工作表数量必须在 1 到 5 之间")
    workbook = Workbook()
    workbook.remove(workbook.active)
    names: set[str] = set()
    cell_count = 0
    try:
        for sheet in sheets:
            if (
                not sheet.name.strip() or sheet.name.casefold() in names
                or any(char in sheet.name for char in "[]:*?/\\")
                or not 1 <= len(sheet.rows) <= MAX_ROWS_PER_SHEET
            ):
                raise ValueError("工作表名称重复、无效或行数超过限制")
            names.add(sheet.name.casefold())
            worksheet = workbook.create_sheet(sheet.name)
            for row_index, row in enumerate(sheet.rows, start=1):
                if len(row) > MAX_COLUMNS:
                    raise ValueError("工作表列数超过限制")
                cell_count += len(row)
                if cell_count > MAX_CELLS:
                    raise ValueError("工作簿单元格超过限制")
                for column, value in enumerate(row, start=1):
                    if isinstance(value, float) and not math.isfinite(value):
                        raise ValueError("工作表数值必须是有限数")
                    if isinstance(value, str) and len(value) > 32767:
                        raise ValueError("单元格文本超过限制")
                    cell = worksheet.cell(row=row_index, column=column, value=value)
                    if isinstance(value, str):
                        cell.data_type = "s"
            worksheet.freeze_panes = "A2" if len(sheet.rows) > 1 else None
        output = io.BytesIO()
        workbook.save(output)
        data = output.getvalue()
        reopened = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            if reopened.sheetnames != [sheet.name for sheet in sheets]:
                raise ValueError("XLSX 生成后校验失败")
        finally:
            reopened.close()
        return data
    except IllegalCharacterError as exc:
        raise ValueError("单元格包含不允许的控制字符") from exc
    finally:
        workbook.close()
